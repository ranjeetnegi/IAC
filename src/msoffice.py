import xlwt
import xlrd
from xlwt import Workbook
from docx import Document
import math

from src.address import Address


class MsOffice:
    def __init__(self):
        self.record_per_sheet = 64_000

    def export_to_MS_word(self, address_list, file_name):
        document = Document()
        no_of_address = len(address_list)
        number_of_rows = math.ceil(no_of_address / 3)
        table = document.add_table(rows=number_of_rows, cols=3)
        table.style = 'Table Grid'
        for i, address in enumerate(address_list):
            row = i % number_of_rows
            col = math.floor(i / number_of_rows)
            print("Updating : %s : %s" % (str(i), address.address))
            cell = table.cell(row, col)
            cell.text = address.address
        document.save(file_name)

    def export_to_MS_Excel(self, address_list, file_name):
        wb = Workbook()
        style_warn = xlwt.easyxf("pattern: pattern solid, fore_colour red;")
        style_alert = xlwt.easyxf("pattern: pattern solid, fore_colour yellow;")
        style_duplicate = xlwt.easyxf("pattern: pattern solid, fore_colour brown;")
        style_repeat = xlwt.easyxf("pattern: pattern solid, fore_colour gray25;")
        row_number = -1
        while row_number < len(address_list):
            row_number = row_number + 1
            if row_number == 0 or row_number % self.record_per_sheet == 0:
                headers_list = ["ADDRESS ORIGINAL","ADDRESS UPDATED","STATE", "DISTRICT", "BLOCK", "PIN", "PHONE", "RE_ORDER", "NAME","DISTRICT_FROM_ADDRESS","STATE_FROM_ADDRESS"
                    ,"DISTRICT_MATCH_COUNT","DIST_MATCHES_PIN_AND_ADDR","STATE_MATCHES_PIN_AND_ADDR","BOOK NAME","BOOK LANG","REPEAT ORDER"]
                wb.add_sheet("Sheet " + str(math.ceil(row_number/self.record_per_sheet)))
                sheet = wb.get_sheet("Sheet " + str(math.ceil(row_number / self.record_per_sheet)))
                self.add_headers_to_sheet(sheet, headers_list)
                continue
            address = address_list[row_number - 1]
            row_index = (row_number % self.record_per_sheet)
            try:
                if address.is_repeat:
                    sheet.write(row_index, 0, address.address_old, style_repeat)
                    sheet.write(row_index, 1, address.address, style_repeat)
                    sheet.write(row_index, 2, address.state, style_repeat)
                    sheet.write(row_index, 3, address.district, style_repeat)
                    sheet.write(row_index, 4, address.block, style_repeat)
                    sheet.write(row_index, 5, address.pin, style_repeat)
                    sheet.write(row_index, 6, address.phone, style_repeat)
                    sheet.write(row_index, 7, "NO", style_repeat)
                    sheet.write(row_index, 8, address.name, style_repeat)
                    sheet.write(row_index, 9, address.district_from_address, style_repeat)
                    sheet.write(row_index, 10, address.state_from_address, style_repeat)
                    sheet.write(row_index, 11, address.occ_count, style_repeat)
                    sheet.write(row_index, 12, address.dist_matches_pin_and_addr, style_repeat)
                    sheet.write(row_index, 13, address.state_matches_pin_and_addr, style_repeat)
                    sheet.write(row_index, 14, address.book_name, style_repeat)
                    sheet.write(row_index, 15, address.book_lang, style_repeat)
                    sheet.write(row_index, 16, "YES", style_repeat)
                    continue

                if address.is_reorder:
                    sheet.write(row_index, 0, address.address_old, style_duplicate)
                    sheet.write(row_index, 1, address.address, style_duplicate)
                    sheet.write(row_index, 2, address.state, style_duplicate)
                    sheet.write(row_index, 3, address.district, style_duplicate)
                    sheet.write(row_index, 4, address.block, style_duplicate)
                    sheet.write(row_index, 5, address.pin, style_duplicate)
                    sheet.write(row_index, 6, address.phone, style_duplicate)
                    sheet.write(row_index, 7, "YES", style_duplicate)
                    sheet.write(row_index, 8, address.name, style_duplicate)
                    sheet.write(row_index, 9, address.district_from_address, style_duplicate)
                    sheet.write(row_index, 10, address.state_from_address, style_duplicate)
                    sheet.write(row_index, 11, address.occ_count, style_duplicate)
                    sheet.write(row_index, 12, address.dist_matches_pin_and_addr, style_duplicate)
                    sheet.write(row_index, 13, address.state_matches_pin_and_addr, style_duplicate)
                    sheet.write(row_index, 14, address.book_name, style_duplicate)
                    sheet.write(row_index, 15, address.book_lang, style_duplicate)
                    sheet.write(row_index, 16, "NO", style_duplicate)
                    continue

                if address.pin is not None and address.phone is not None:
                    sheet.write(row_index, 0, address.address_old)
                    sheet.write(row_index, 1, address.address)
                    sheet.write(row_index, 2, address.state)
                    sheet.write(row_index, 3, address.district)
                    sheet.write(row_index, 4, address.block)
                    sheet.write(row_index, 5, address.pin)
                    sheet.write(row_index, 6, address.phone)
                    sheet.write(row_index, 7, "NO")
                    sheet.write(row_index, 8, address.name)
                    sheet.write(row_index, 9, address.district_from_address)
                    sheet.write(row_index, 10, address.state_from_address)
                    sheet.write(row_index, 11, address.occ_count)
                    sheet.write(row_index, 12, address.dist_matches_pin_and_addr)
                    sheet.write(row_index, 13, address.state_matches_pin_and_addr)
                    sheet.write(row_index, 14, address.book_name)
                    sheet.write(row_index, 15, address.book_lang)
                    sheet.write(row_index, 16, "NO")
                elif address.phone is not None:
                    sheet.write(row_index, 0, address.address_old, style_alert)
                    sheet.write(row_index, 1, address.address, style_alert)
                    sheet.write(row_index, 2, address.state, style_alert)
                    sheet.write(row_index, 3, address.district, style_alert)
                    sheet.write(row_index, 4, address.block, style_alert)
                    sheet.write(row_index, 5, address.pin, style_alert)
                    sheet.write(row_index, 6, address.phone, style_alert)
                    sheet.write(row_index, 7, "NO", style_alert)
                    sheet.write(row_index, 8, address.name,style_alert)
                    sheet.write(row_index, 9, address.district_from_address, style_alert)
                    sheet.write(row_index, 10, address.state_from_address, style_alert)
                    sheet.write(row_index, 11, address.occ_count, style_alert)
                    sheet.write(row_index, 12, address.dist_matches_pin_and_addr,style_alert)
                    sheet.write(row_index, 13, address.state_matches_pin_and_addr,style_alert)
                    sheet.write(row_index, 14, address.book_name, style_alert)
                    sheet.write(row_index, 15, address.book_lang, style_alert)
                    sheet.write(row_index, 16, "NO", style_alert)
                else:
                    sheet.write(row_index, 0, address.address_old, style_warn)
                    sheet.write(row_index, 1, address.address, style_warn)
                    sheet.write(row_index, 2, address.state, style_warn)
                    sheet.write(row_index, 3, address.district, style_warn)
                    sheet.write(row_index, 4, address.block, style_warn)
                    sheet.write(row_index, 5, address.pin, style_warn)
                    sheet.write(row_index, 6, address.phone, style_warn)
                    sheet.write(row_index, 7, "NO", style_warn)
                    sheet.write(row_index, 8, address.name, style_warn)
                    sheet.write(row_index, 9, address.district_from_address, style_warn)
                    sheet.write(row_index, 10, address.state_from_address, style_warn)
                    sheet.write(row_index, 11, address.occ_count, style_warn)
                    sheet.write(row_index, 12, address.dist_matches_pin_and_addr,style_warn)
                    sheet.write(row_index, 13, address.state_matches_pin_and_addr, style_warn)
                    sheet.write(row_index, 14, address.book_name, style_warn)
                    sheet.write(row_index, 15, address.book_lang, style_warn)
                    sheet.write(row_index, 16, "NO", style_warn)
            except:
                address.print_attributes()
        wb.save(file_name)

    def add_headers_to_sheet(self, worksheet, headers_list):
        style_bold_black_color = xlwt.easyxf("align:wrap on; font: bold on, color-index black")
        for i, header_name in enumerate(headers_list):
            worksheet.write(0, i, header_name, style_bold_black_color)

    def update_sheet_cell(self, sheet, row, col, data, style):
        sheet.write(row, col, data, style)
        return

    def import_from_Excel_sheet(self, file_name):
        wb = xlrd.open_workbook(file_name)
        sheets = wb.sheet_names()
        address_list = []
        for i in sheets:
            sheet = wb.sheet_by_name(str(i))
            rows = sheet.nrows
            for row_index in range(1, rows):
                address_text = sheet.cell_value(row_index, 0)
                state = sheet.cell_value(row_index, 1)
                if len(state) == 0:
                    state = None
                district = sheet.cell_value(row_index, 2)
                if len(district) == 0:
                    district = None
                block = sheet.cell_value(row_index, 3)
                if len(block) == 0:
                    block = None
                pin = sheet.cell_value(row_index, 4)
                if len(pin) == 0:
                    pin = None
                phone = sheet.cell_value(row_index, 5)
                if len(phone) == 0:
                    phone = None
                re_order_text = sheet.cell_value(row_index, 6)
                re_order = None
                if re_order_text == "NO":
                    re_order = False
                elif re_order_text == "YES":
                    re_order = True
                name = sheet.cell_value(row_index, 7)
                dfa = sheet.cell_value(row_index, 8)
                sfa = sheet.cell_value(row_index, 9)
                oc = sheet.cell_value(row_index, 10)
                dmpaa = sheet.cell_value(row_index,11)
                smpaa = sheet.cell_value(row_index,12)
                address_obj = Address(address_text, state, district, block, pin, phone, re_order, name, sfa, dfa, oc,
                                      dmpaa, smpaa)
                address_list.append(address_obj)
        return address_list
