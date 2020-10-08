import xlwt
import xlrd
from xlwt import Workbook
from docx import Document
from datetime import date, time, datetime
import math

from src.address import Address


class MsOffice:
    def __init__(self):
        pass

    def export_to_MS_word(self, address_list, file_name):
        document = Document()
        no_of_address = len(address_list)
        number_of_rows = math.ceil(no_of_address / 3)
        table = document.add_table(rows=number_of_rows, cols=3)
        table.style = 'TableGrid'
        for i, address in enumerate(address_list):
            row = i % number_of_rows
            col = math.floor(i / number_of_rows)
            print("Updating row : " + str(row) + " col : " + str(col) + " address: "+ address.address)
            cell = table.cell(row, col)
            cell.text = address.address
        document.save(file_name)

    def export_to_MS_Excel(self, address_list, file_name):
        wb = Workbook()
        sheet1 = wb.add_sheet("Sheet 1")
        style_warn = xlwt.easyxf("pattern: pattern solid, fore_colour red;")
        style_alert = xlwt.easyxf("pattern: pattern solid, fore_colour yellow;")
        row_number = -1
        while row_number < len(address_list):
            row_number = row_number + 1
            if row_number == 0:
                headers_list = ["ADDRESS", "STATE", "DISTRICT", "BLOCK", "PIN", "PHONE"]
                self.add_headers_to_sheet(sheet1, headers_list)
                continue
            address = address_list[row_number - 1]
            try:
                if address.pin != None and address.phone != None:
                    sheet1.write(row_number, 0, address.address)
                    sheet1.write(row_number, 1, address.state)
                    sheet1.write(row_number, 2, address.district)
                    sheet1.write(row_number, 3, address.block)
                    sheet1.write(row_number, 4, address.pin)
                    sheet1.write(row_number, 5, address.phone)
                elif address.phone != None:
                    sheet1.write(row_number, 0, address.address, style_alert)
                    sheet1.write(row_number, 1, address.state, style_alert)
                    sheet1.write(row_number, 2, address.district, style_alert)
                    sheet1.write(row_number, 3, address.block, style_alert)
                    sheet1.write(row_number, 4, address.pin, style_alert)
                    sheet1.write(row_number, 5, address.phone, style_alert)
                else:
                    sheet1.write(row_number, 0, address.address, style_warn)
                    sheet1.write(row_number, 1, address.state, style_warn)
                    sheet1.write(row_number, 2, address.district, style_warn)
                    sheet1.write(row_number, 3, address.block, style_warn)
                    sheet1.write(row_number, 4, address.pin, style_warn)
                    sheet1.write(row_number, 5, address.phone, style_warn)
            except:
                address.print_attributes()
        wb.save(file_name)

    def add_headers_to_sheet(self, worksheet, headers_list):
        style_bold_black_color = xlwt.easyxf(
            "align:wrap on; font: bold on, color-index black"
        )
        for i, header_name in enumerate(headers_list):
            worksheet.write(0, i, header_name, style_bold_black_color)

    def import_from_Excel_sheet(self, file_name, sheet_number, address_col):
        wb = xlrd.open_workbook(file_name)
        sheet = wb.sheet_by_index(sheet_number)
        address_list = []
        start_row = 1
        while True:
            address_text = sheet.cell_value(start_row, address_col)
            if address_text == None or len(address_text) == 0:
                if len(sheet.cell_value(start_row + 1, address_col)) == 0:
                    break
            address_obj = Address(address_text, None, None, None, None, None)
            address_list.append(address_obj)
        return address_list